"""Unit tests for spreadsheet extraction + row-aware chunking.

Builds real .xlsx workbooks in memory with openpyxl (skipped where it is not
installed), so no service and no MarkItDown are needed:

    python rag_service/spreadsheet_test.py
"""
import datetime
import io
import os
import sys
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(os.path.dirname(HERE), "poc"))

import extract  # noqa: E402
from chunking import chunk_rows  # noqa: E402

try:
    import openpyxl
except Exception:  # noqa: BLE001
    openpyxl = None


def _workbook(sheets):
    """sheets: list of (title, rows) where rows is a list of lists (None = empty)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets:
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@unittest.skipIf(openpyxl is None, "openpyxl not installed")
class WorkbookExtractionTest(unittest.TestCase):
    def test_each_sheet_is_a_page_named_after_the_sheet(self):
        raw = _workbook([
            ("Budget", [["Item", "Amount"], ["Rent", 1200]]),
            ("Notes", [["Remember to renew"]]),
        ])
        pages = extract.extract_pages(raw, "/drive/fin/plan.xlsx")
        self.assertEqual([(1, "Budget"), (2, "Notes")], [(p[0], p[1]) for p in pages])
        self.assertTrue(pages[0][2].startswith("Sheet: Budget\n"))

    def test_empty_cells_are_absent_not_nan(self):
        raw = _workbook([("S", [
            [None, None, "Title only"],
            [None, None, None],
            ["Name", "Role", "Note"],
            ["Anna", None, "on leave"],
            [None, None, None],
            ["Bela", "Dev", None],
        ])])
        text = extract.extract_pages(raw, "/d/a.xlsx")[0][2]
        self.assertNotIn("NaN", text)
        self.assertNotIn("None", text)
        self.assertNotIn("Unnamed", text)
        self.assertIn("Columns: Name | Role | Note", text)
        self.assertIn("Name: Anna; Note: on leave", text)
        self.assertIn("Name: Bela; Role: Dev", text)
        # The banner above the header stays as prose, not as a record.
        self.assertIn("\nTitle only\n", text)

    def test_header_labels_columns_even_when_header_has_gaps(self):
        raw = _workbook([("S", [
            [None, "Task", None, "Owner"],
            [None, "Patch servers", "urgent", "Ops"],
        ])])
        text = extract.extract_pages(raw, "/d/a.xlsx")[0][2]
        self.assertIn("Task: Patch servers; C: urgent; Owner: Ops", text)

    def test_numbers_dates_and_percentages_are_readable(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(["Metric", "Value", "When", "Share", "Flag"])
        ws.append(["Revenue", 14.7, datetime.datetime(2026, 6, 30), 0.256, True])
        ws.append(["Sites", 41.0, datetime.datetime(2026, 6, 30, 14, 5), 1 / 3, False])
        ws["D2"].number_format = "0.0%"
        ws["D3"].number_format = "0%"
        buf = io.BytesIO()
        wb.save(buf)
        text = extract.extract_pages(buf.getvalue(), "/d/a.xlsx")[0][2]
        self.assertIn("Value: 14.7", text)
        self.assertIn("Value: 41;", text)          # 41.0 -> 41
        self.assertIn("When: 2026-06-30;", text)   # midnight -> date only
        self.assertIn("When: 2026-06-30 14:05", text)
        self.assertIn("Share: 25.6%", text)
        self.assertIn("Share: 33.33333333%", text)
        self.assertIn("Flag: TRUE", text)
        self.assertIn("Flag: FALSE", text)

    def test_sheet_without_header_is_rendered_as_plain_rows(self):
        raw = _workbook([("S", [[1, 2, 3], [4, 5, 6]])])
        text = extract.extract_pages(raw, "/d/a.xlsx")[0][2]
        self.assertEqual("Sheet: S\n1 | 2 | 3\n4 | 5 | 6", text)
        self.assertEqual(1, extract.header_lines("/d/a.xlsx", text))

    def test_empty_sheets_are_skipped_and_empty_workbook_yields_no_pages(self):
        raw = _workbook([("Empty", []), ("Data", [["a", "b"], ["x", "y"]])])
        pages = extract.extract_pages(raw, "/d/a.xlsx")
        self.assertEqual([(2, "Data")], [(p[0], p[1]) for p in pages])
        self.assertEqual([], extract.extract_pages(_workbook([("E", [])]), "/d/a.xlsx"))

    def test_row_cap_is_reported_not_silent(self):
        saved = extract.SHEET_MAX_ROWS
        extract.SHEET_MAX_ROWS = 3
        try:
            raw = _workbook([("S", [["k", "v"]] + [["r%d" % i, i] for i in range(10)])])
            text = extract.extract_pages(raw, "/d/a.xlsx")[0][2]
        finally:
            extract.SHEET_MAX_ROWS = saved
        self.assertIn("(8 further rows of this sheet were not indexed)", text)
        self.assertIn("k: r1; v: 1", text)
        self.assertNotIn("r5", text)

    def test_corrupt_workbook_raises_extraction_error(self):
        with self.assertRaises(extract.ExtractionError):
            extract.extract_pages(b"PK\x03\x04 definitely not a workbook", "/d/a.xlsx")

    def test_xlsm_is_supported_and_xls_still_goes_through_markitdown(self):
        self.assertTrue(extract.is_supported("/d/macro.xlsm"))
        self.assertTrue(extract.is_tabular("/d/macro.xlsm"))
        self.assertTrue(extract.extractor_available("/d/macro.xlsm"))
        self.assertFalse(extract.is_tabular("/d/legacy.xls"))
        self.assertTrue(extract.needs_markitdown("/d/legacy.xls"))
        self.assertFalse(extract.needs_markitdown("/d/new.xlsx"))


class ChunkRowsTest(unittest.TestCase):
    def test_header_is_repeated_and_rows_never_split(self):
        rows = ["Name: n%d; Note: word word word word" % i for i in range(20)]
        text = "Sheet: S\nColumns: Name | Note\n" + "\n".join(rows)
        chunks = chunk_rows(text, size=30, overlap=0, header_lines=2)
        self.assertGreater(len(chunks), 1)
        for c in chunks:
            lines = c.split("\n")
            self.assertEqual(["Sheet: S", "Columns: Name | Note"], lines[:2])
            for body in lines[2:]:
                self.assertIn(body, rows)
        seen = [ln for c in chunks for ln in c.split("\n")[2:]]
        self.assertEqual(rows, seen)  # every row exactly once with overlap=0

    def test_overlap_carries_whole_trailing_rows(self):
        rows = ["r%d a b c" % i for i in range(8)]  # 4 words each
        chunks = chunk_rows("\n".join(rows), size=12, overlap=4, header_lines=0)
        self.assertEqual(["r0 a b c", "r1 a b c", "r2 a b c"], chunks[0].split("\n"))
        self.assertEqual("r2 a b c", chunks[1].split("\n")[0])  # carried row

    def test_oversized_row_becomes_its_own_chunk_without_duplicating_tail(self):
        big = " ".join(["w"] * 50)
        chunks = chunk_rows("\n".join(["a b", "c d", big, "e f"]), size=10, overlap=4)
        self.assertEqual(["a b\nc d", big, "e f"], chunks)

    def test_header_only_and_empty(self):
        self.assertEqual(["Sheet: S"], chunk_rows("Sheet: S\n", 10, 2, header_lines=1))
        self.assertEqual([], chunk_rows("   \n  ", 10, 2))

    def test_csv_repeats_first_line(self):
        text = "id,name,amount\n1,Anna,10\n2,Bela,20\n3,Cili,30"
        self.assertEqual(1, extract.header_lines("/d/x.csv", text))
        chunks = chunk_rows(text, size=3, overlap=0, header_lines=1)
        self.assertEqual(["id,name,amount\n1,Anna,10\n2,Bela,20", "id,name,amount\n3,Cili,30"], chunks)


if __name__ == "__main__":
    unittest.main()
