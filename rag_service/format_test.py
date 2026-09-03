"""Real-format extraction test for the RAG service.

Designed to run INSIDE the rag_service container (where MarkItDown and its
format handlers are installed), hitting the already-running service on
127.0.0.1:8090:

    docker exec vivesec-rag python rag_service/format_test.py

It generates genuine office documents (DOCX/PPTX/XLSX via python-docx /
python-pptx / openpyxl, all pulled by markitdown[all]), an HTML and a CSV, and
a minimal hand-built PDF, each carrying a unique marker phrase. Every file is
pushed through the single-step content endpoint and then the marker is searched
back, proving end-to-end extraction + embedding + retrieval on real formats.
"""
import base64
import json
import os
import sys
import urllib.error
import urllib.request

PORT = int(os.environ.get("RAG_PORT", "8090"))
BASE = "http://127.0.0.1:%d" % PORT
CORPUS = "format-test"

_failures = []


def _req(method, path, body=None):
    data = json.dumps(body).encode("utf-8") if body is not None else None
    headers = {"Content-Type": "application/json"} if data else {}
    req = urllib.request.Request(BASE + path, data=data, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req, timeout=120) as resp:
            return resp.status, json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        return e.code, json.loads(e.read().decode("utf-8"))


def check(name, cond, detail=""):
    print("[%s] %s%s" % ("PASS" if cond else "FAIL", name, (" -- " + detail) if detail else ""))
    if not cond:
        _failures.append(name)


# ---- format generators -----------------------------------------------------
def gen_docx(marker):
    from docx import Document
    from io import BytesIO
    doc = Document()
    doc.add_heading("Confidential Memo", level=1)
    doc.add_paragraph("This document discusses the annual strategy review.")
    doc.add_paragraph("Special marker: %s appears in the body text." % marker)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def gen_pptx(marker):
    from pptx import Presentation
    from io import BytesIO
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Roadmap"
    slide.placeholders[1].text = "Phase two milestone. Marker %s on this slide." % marker
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def gen_xlsx(marker):
    from openpyxl import Workbook
    from io import BytesIO
    wb = Workbook()
    ws = wb.active
    ws.append(["Region", "Revenue", "Note"])
    ws.append(["EMEA", 42000000, "marker %s" % marker])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def gen_html(marker):
    html = (
        "<html><head><title>Report</title></head><body>"
        "<h1>Quarterly Report</h1>"
        "<p>The infrastructure audit passed. Marker %s is here.</p>"
        "</body></html>" % marker
    )
    return html.encode("utf-8")


def gen_csv(marker):
    return ("col_a,col_b\nvalue,marker %s\n" % marker).encode("utf-8")


def gen_pdf(marker):
    """Minimal single-page PDF with a text string (xref offsets computed)."""
    text = ("Audit summary. Marker %s in the PDF." % marker).replace("(", "").replace(")", "")
    stream = b"BT /F1 18 Tf 72 720 Td (" + text.encode("latin-1") + b") Tj ET"
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += str(i).encode() + b" 0 obj\n" + body + b"\nendobj\n"
    xref_pos = len(out)
    out += b"xref\n0 " + str(len(objs) + 1).encode() + b"\n"
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += ("%010d 00000 n \n" % off).encode()
    out += b"trailer\n<< /Size " + str(len(objs) + 1).encode() + b" /Root 1 0 R >>\n"
    out += b"startxref\n" + str(xref_pos).encode() + b"\n%%EOF"
    return out


CASES = [
    ("memo.docx", "DOCX", gen_docx, "zeta-docx-7731"),
    ("deck.pptx", "PPTX", gen_pptx, "zeta-pptx-4412"),
    ("sheet.xlsx", "XLSX", gen_xlsx, "zeta-xlsx-9920"),
    ("report.html", "HTML", gen_html, "zeta-html-3185"),
    ("data.csv", "CSV", gen_csv, "zeta-csv-5567"),
    ("audit.pdf", "PDF", gen_pdf, "zeta-pdf-8043"),
]


def main():
    status, health = _req("GET", "/health")
    emb = health.get("embedding", {})
    print("embedding: %s, vector_backend: %s" % (json.dumps(emb), health.get("vector_backend")))

    for fname, label, gen, marker in CASES:
        try:
            raw = gen(marker)
        except Exception as e:  # noqa: BLE001
            check("%s generate" % label, False, str(e))
            continue
        path = "/format-test/%s" % fname
        status, res = _req("POST", "/index/upsert/file/content", {
            "corpus_id": CORPUS, "path": path, "title": fname,
            "content_b64": base64.b64encode(raw).decode("ascii"),
        })
        ok = status == 200 and res.get("ok") and (res.get("chunks") or 0) >= 1
        check("%s ingest -> chunks>=1" % label, ok, json.dumps(res))
        if not ok:
            continue
        # retrieve the marker back
        status, s = _req("POST", "/rag/search_context", {
            "corpus_id": CORPUS, "question": "marker %s" % marker, "top_k": 3,
        })
        contexts = s.get("contexts", [])
        hit = any(marker in (c.get("text") or "") for c in contexts)
        check("%s marker retrievable" % label, hit,
              "top_path=%s" % (contexts[0].get("source_path") if contexts else "none"))

    print("\n%d check(s) failed: %s" % (len(_failures), _failures) if _failures else "\nALL FORMAT CHECKS PASSED")
    sys.exit(1 if _failures else 0)


if __name__ == "__main__":
    main()
