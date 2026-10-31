"""XLSX writer and investigation export (E26b, BV 1.5 / 2.5)."""
import io
import os
import unittest
import zipfile
import xml.etree.ElementTree as ET
from unittest.mock import patch

import investigation_export as ie
import orchestrator
import xlsx
from orchestrator_test import InvestigationGraph

NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
CONFIG = {"relations": ["RELATIVE_OF"], "via": ["INVOLVED_IN"], "branch": 10, "max_rows": 80}
DEFAULT = {"columns": list(ie.COLUMNS), "sheets": "single"}

try:
    import openpyxl  # noqa: F401 -- dev-only independent reader; the adapter never imports it
    HAVE_OPENPYXL = True
except ImportError:
    HAVE_OPENPYXL = False


def parts(content):
    z = zipfile.ZipFile(io.BytesIO(content))
    return {name: z.read(name).decode("utf-8") for name in z.namelist()}


def sheet_rows(xml):
    """The cell values of a worksheet, row by row (text or number strings)."""
    out = []
    for row in ET.fromstring(xml).find("m:sheetData", NS):
        values = {}
        for c in row:
            col = "".join(ch for ch in c.get("r") if ch.isalpha())
            text = c.find("m:is/m:t", NS)
            values[col] = text.text if text is not None else c.find("m:v", NS).text
        width = max((ord(k[-1]) - 64 + 26 * (len(k) - 1) for k in values), default=0)
        out.append([values.get(xlsx.column_letter(i)) for i in range(width)])
    return out


def run_investigation(graph=None, **kw):
    graph = graph or InvestigationGraph()
    run = orchestrator.Run(graph, {"corpus_ids": ["c"], "access": None},
                           lambda p: "/hr/" not in p, "Hungarian", investigate=CONFIG, **kw)
    return ie.investigate(run, "Sas Tibor", 2), graph


class XlsxWriterTest(unittest.TestCase):
    def book(self):
        return xlsx.workbook([
            {"name": "Adatok: [1/2]", "rows": [["Név", "Szint"], ["=1+1", 2],
                                                ["a\x00b\x1fc", 1.5], ["x" * 40000, None]]},
            {"name": "Adatok: [1/2]", "rows": [["k", "v"]], "filter": False}], title="T & T")

    def test_package_parts_are_well_formed_and_linked(self):
        files = parts(self.book())
        for name in ("[Content_Types].xml", "_rels/.rels", "xl/workbook.xml",
                     "xl/_rels/workbook.xml.rels", "xl/styles.xml", "docProps/core.xml",
                     "docProps/app.xml", "xl/worksheets/sheet1.xml",
                     "xl/worksheets/sheet2.xml"):
            self.assertIn(name, files)
            ET.fromstring(files[name])  # well-formed
        types = files["[Content_Types].xml"]
        self.assertIn('PartName="/xl/worksheets/sheet2.xml"', types)
        self.assertIn("spreadsheetml.sheet.main+xml", types)
        book = ET.fromstring(files["xl/workbook.xml"])
        names = [s.get("name") for s in book.find("m:sheets", NS)]
        # Excel: no []:*?/\, at most 31 characters, unique regardless of case.
        self.assertEqual(["Adatok_ _1_2_", "Adatok_ _1_2_ (2)"], names)
        self.assertIn("worksheets/sheet2.xml", files["xl/_rels/workbook.xml.rels"])
        self.assertIn("&amp;", files["docProps/core.xml"])

    def test_cells_are_text_or_numbers_never_formulas(self):
        files = parts(self.book())
        sheet = files["xl/worksheets/sheet1.xml"]
        self.assertNotIn("<f>", sheet)
        rows = sheet_rows(sheet)
        self.assertEqual(["Név", "Szint"], rows[0])
        self.assertEqual(["=1+1", "2"], rows[1])
        self.assertEqual(["abc", "1.5"], rows[2])
        self.assertEqual(xlsx.MAX_CELL, len(rows[3][0]))
        root = ET.fromstring(sheet)
        self.assertEqual("1", root.find("m:sheetData/m:row/m:c", NS).get("s"))  # bold header
        self.assertEqual("frozen", root.find("m:sheetViews/m:sheetView/m:pane", NS).get("state"))
        self.assertEqual("A1:B4", root.find("m:autoFilter", NS).get("ref"))
        self.assertIsNone(ET.fromstring(files["xl/worksheets/sheet2.xml"]).find(
            "m:autoFilter", NS))
        self.assertIn("_xlnm._FilterDatabase", files["xl/workbook.xml"])

    def test_column_letters(self):
        self.assertEqual(["A", "Z", "AA", "AZ", "BA"],
                         [xlsx.column_letter(i) for i in (0, 25, 26, 51, 52)])

    @unittest.skipUnless(HAVE_OPENPYXL, "openpyxl (dev only) not installed")
    def test_an_independent_reader_opens_it(self):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(self.book()))
        self.assertEqual(["Adatok_ _1_2_", "Adatok_ _1_2_ (2)"], wb.sheetnames)
        ws = wb.worksheets[0]
        self.assertEqual(("Név", "Szint"), (ws["A1"].value, ws["B1"].value))
        self.assertEqual(("=1+1", 2, 1.5), (ws["A2"].value, ws["B2"].value, ws["B3"].value))
        self.assertEqual("s", ws["A2"].data_type)
        self.assertTrue(ws["A1"].font.b)
        self.assertEqual("A2", ws.freeze_panes)
        self.assertEqual("A1:B4", ws.auto_filter.ref)


class InvestigationExportTest(unittest.TestCase):
    def build(self, inv, lang="hu", conf=None, **meta):
        content, name, title = ie.build(inv, lang, conf or DEFAULT,
                                        dict({"config": CONFIG, "now": 1793000000}, **meta))
        return parts(content), name, title, content

    def test_columns_in_order_with_the_visible_rows(self):
        inv, _ = run_investigation()
        files, name, title, _ = self.build(inv, user="lars", drives=["ops"], case_ref="BV-7",
                                           legal_basis="Bv. tv. 1. §", audit_id="abc123")
        rows = sheet_rows(files["xl/worksheets/sheet1.xml"])
        self.assertEqual(["Szint", "Név", "Szerep", "Nyilvántartási szám", "Azonosító szám",
                          "Telefon", "Lakcím", "E-mail", "Egyéb azonosítók",
                          "Kapcsolódás útja", "További szál", "Forrás", "Oldal",
                          "Összevont rekordok"], rows[0])
        self.assertEqual(["Sas Tibor", "Kende Anna", "Major Ede", "Izsó Csaba"],
                         [r[1] for r in rows[1:]])
        self.assertEqual(["0", "Sas Tibor", "Fogvatartott", "NY-100"], rows[1][:4])
        self.assertEqual("kiindulópont", rows[1][9])
        kende = rows[2]
        self.assertEqual(("+36 30 111 1111", "Fő utca 5.", None),
                         (kende[5], kende[6], kende[7]))
        self.assertEqual("Sas Tibor \u2014[Rokona]\u2014 Kende Anna", kende[9])
        self.assertEqual("/storage/drives/ops/a.txt", kende[11])
        self.assertIn("Érintett: verekedés (2025-03-01)", rows[3][9])
        self.assertTrue(name.startswith("vizsgalat_Sas_Tibor_"))
        self.assertEqual("Vizsgálati tábla \u2013 Sas Tibor", title)
        info = dict((r + [None])[:2] for r in sheet_rows(files["xl/worksheets/sheet2.xml"]))
        self.assertEqual(("Sas Tibor", "2", "4", "ops", "BV-7", "Bv. tv. 1. §", "abc123", "lars"),
                         (info["Kiindulópont"], info["Kért szintek"], info["Sorok"],
                          info["Hatókör (meghajtók)"], info["Ügyszám"], info["Jogalap"],
                          info["A válasz audit-azonosítója"], info["Kérte"]))
        self.assertEqual("Ezekre a szintekre teljes.", info["Teljesség"])
        self.assertIn("nem a modell válaszából", info["Megjegyzés"])

    def test_what_the_adapter_hides_is_not_in_the_file(self):
        inv, _ = run_investigation()
        _, _, _, content = self.build(inv)
        blob = "".join(parts(content).values())
        # Kiss Bori rests on a hidden document only; Illés Dénes is reached only
        # through her; Kende Anna's e-mail and Sas Tibor's second number are hidden too.
        for hidden in ("Kiss Bori", "Illés Dénes", "kende@example.hu", "NY-999", "/hr/"):
            self.assertNotIn(hidden, blob)

    def test_nothing_visible_gives_no_table(self):
        graph = InvestigationGraph()
        run = orchestrator.Run(graph, {"corpus_ids": ["c"], "access": None},
                               lambda p: False, "Hungarian", investigate=CONFIG)
        self.assertIsNone(ie.investigate(run, "Sas Tibor", 2))

    def test_one_sheet_per_level_and_a_cut_table(self):
        inv, _ = run_investigation(InvestigationGraph(truncated=True))
        files, _, _, content = self.build(inv, "en", {"columns": ["name", "level"],
                                                      "sheets": "per_level"})
        names = [s.get("name") for s in ET.fromstring(files["xl/workbook.xml"]).find(
            "m:sheets", NS)]
        self.assertEqual(["Starting point", "Level 1", "Level 2", "About"], names)
        self.assertEqual([["Name", "Level"], ["Major Ede", "1"], ["Kende Anna", "1"]],
                         sorted(sheet_rows(files["xl/worksheets/sheet2.xml"]), reverse=True))
        info = dict((r + [None])[:2] for r in sheet_rows(files["xl/worksheets/sheet4.xml"]))
        self.assertIn("Cut: at most 10 new people from one person, 80 rows", info["Coverage"])
        if HAVE_OPENPYXL:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content))
            self.assertEqual([("Name", "Level"), ("Izsó Csaba", 2)],
                             list(wb["Level 2"].iter_rows(values_only=True)))

    def test_settings_fall_back_to_the_safe_layout(self):
        with patch("sys.stderr"):
            conf = ie.settings({"ADAPTER_INVESTIGATION_XLSX_COLUMNS": "level, phone, bogus, phone",
                                "ADAPTER_INVESTIGATION_XLSX_SHEETS": "tabs"})
        # The name and the source of a row are always there.
        self.assertEqual({"columns": ["level", "name", "phone", "source"], "sheets": "single"},
                         conf)
        self.assertEqual(DEFAULT, ie.settings({}))
        with patch("sys.stderr"):
            self.assertEqual(list(ie.COLUMNS), ie.settings(
                {"ADAPTER_INVESTIGATION_XLSX_COLUMNS": "nope"})["columns"])
        self.assertEqual("per_level", ie.settings(
            {"ADAPTER_INVESTIGATION_XLSX_SHEETS": " Per_Level "})["sheets"])

    def test_request_validation(self):
        self.assertEqual(("Sas Tibor", 2), ie.parse_request({"start": " Sas Tibor "}))
        self.assertEqual(("NY-1", 4), ie.parse_request({"start": "NY-1", "levels": "4"}))
        for bad in ({}, {"start": ""}, {"start": {"name": "x"}}, {"start": "x" * 121},
                    {"start": "x", "levels": 5}, {"start": "x", "levels": 0},
                    {"start": "x", "levels": True}, {"start": "x", "levels": "two"}):
            with self.assertRaises(ValueError, msg=bad):
                ie.parse_request(bad)
        self.assertEqual(("hu", "en", "en"), (ie.lang_code("HU"), ie.lang_code("xx"),
                                              ie.lang_code(None)))


class StoredFormatTest(unittest.TestCase):
    def test_xlsx_is_stored_but_not_rendered(self):
        import docgen
        self.assertEqual("xlsx", docgen.stored_format(".XLSX"))
        self.assertEqual("docx", docgen.stored_format("docx"))
        self.assertIsNone(docgen.normalize_format("xlsx"))
        self.assertIsNone(docgen.stored_format("exe"))
        self.assertEqual("a.xlsx", docgen.with_extension("a.xlsx", "xlsx"))
        self.assertEqual(xlsx.CONTENT_TYPE, docgen.content_type("xlsx"))


if __name__ == "__main__":
    unittest.main()
