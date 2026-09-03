"""Minimal, dependency-light document writers for the demo corpus.

Only `openpyxl` is an external dependency (real .xlsx). PDF and DOCX are
written by hand so the generator runs anywhere with plain Python 3.10+.

Formats produced:
    write_text   -> .md / .txt / .csv (UTF-8, LF)
    write_xlsx   -> real spreadsheet (openpyxl)
    write_docx   -> minimal but valid OOXML word document
    write_pdf    -> multi-page text PDF (Helvetica, WinAnsi)
    write_scanned_pdf -> image-only PDF with NO text layer (OCR test case)
"""
from __future__ import annotations

import io
import os
import random
import textwrap
import zipfile
import zlib
from xml.sax.saxutils import escape

# --------------------------------------------------------------------------
# plain text
# --------------------------------------------------------------------------


def write_text(path: str, body: str) -> int:
    os.makedirs(os.path.dirname(path), exist_ok=True)
    data = body if body.endswith("\n") else body + "\n"
    with io.open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write(data)
    return os.path.getsize(path)


# --------------------------------------------------------------------------
# xlsx
# --------------------------------------------------------------------------


def write_xlsx(path: str, sheets: dict[str, list[list]]) -> int:
    """sheets: {sheet_name: [[row cells], ...]}; the first row is treated as header."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name[:31])
        for row in rows:
            # Blank cells are written as a dash. openpyxl stores None and "" as
            # a truly empty cell, and the spreadsheet-to-text extractors used
            # for indexing render an empty cell as "NaN" — which then leaks into
            # generated answers ("it states NaN"). A dash reads correctly both
            # in the spreadsheet and in the extracted text.
            ws.append(["-" if c is None or c == "" else c for c in row])
        if rows:
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(vertical="center")
            for idx, _ in enumerate(rows[0], start=1):
                letter = ws.cell(row=1, column=idx).column_letter
                width = max(
                    (len(str(r[idx - 1])) for r in rows if len(r) >= idx and r[idx - 1] is not None),
                    default=10,
                )
                ws.column_dimensions[letter].width = min(max(width + 2, 10), 48)
            ws.freeze_panes = "A2"
    wb.save(path)
    return os.path.getsize(path)


# --------------------------------------------------------------------------
# docx (hand-rolled minimal OOXML)
# --------------------------------------------------------------------------

_DOCX_CONTENT_TYPES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>
<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>
</Types>"""

_DOCX_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>
</Relationships>"""

_DOCX_DOC_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>
</Relationships>"""

_DOCX_STYLES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:style w:type="paragraph" w:styleId="Normal" w:default="1"><w:name w:val="Normal"/></w:style>
<w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/><w:basedOn w:val="Normal"/>
<w:pPr><w:outlineLvl w:val="0"/></w:pPr><w:rPr><w:b/><w:sz w:val="32"/></w:rPr></w:style>
<w:style w:type="paragraph" w:styleId="Heading2"><w:name w:val="heading 2"/><w:basedOn w:val="Normal"/>
<w:pPr><w:outlineLvl w:val="1"/></w:pPr><w:rPr><w:b/><w:sz w:val="26"/></w:rPr></w:style>
</w:styles>"""


def _docx_core(title: str) -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
        'xmlns:dc="http://purl.org/dc/elements/1.1/">'
        "<dc:title>%s</dc:title><dc:creator>Voltara Energy Group</dc:creator>"
        "</cp:coreProperties>" % escape(title)
    )


def _docx_paragraph(text: str, style: str | None = None) -> str:
    ppr = '<w:pPr><w:pStyle w:val="%s"/></w:pPr>' % style if style else ""
    if not text:
        return "<w:p>%s</w:p>" % ppr
    return '<w:p>%s<w:r><w:t xml:space="preserve">%s</w:t></w:r></w:p>' % (ppr, escape(text))


def write_docx(path: str, title: str, blocks: list[tuple[str, str]]) -> int:
    """blocks: list of (kind, text) where kind is 'h1' | 'h2' | 'p' | ''."""
    os.makedirs(os.path.dirname(path), exist_ok=True)
    style_map = {"h1": "Heading1", "h2": "Heading2"}
    body = [_docx_paragraph(title, "Heading1")]
    for kind, text in blocks:
        body.append(_docx_paragraph(text, style_map.get(kind)))
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        "<w:body>%s<w:sectPr><w:pgSz w:w=\"11906\" w:h=\"16838\"/></w:sectPr></w:body></w:document>"
        % "".join(body)
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", _DOCX_CONTENT_TYPES)
        z.writestr("_rels/.rels", _DOCX_RELS)
        z.writestr("docProps/core.xml", _docx_core(title))
        z.writestr("word/_rels/document.xml.rels", _DOCX_DOC_RELS)
        z.writestr("word/styles.xml", _DOCX_STYLES)
        z.writestr("word/document.xml", document)
    return os.path.getsize(path)


# --------------------------------------------------------------------------
# pdf
# --------------------------------------------------------------------------

_PAGE_W, _PAGE_H = 595, 842  # A4 in points
_MARGIN_X, _MARGIN_TOP = 56, 786
_LEADING = 14
_LINES_PER_PAGE = 50
_WRAP = 92


# Characters WinAnsi (cp1252) cannot represent but the corpus needs. They are
# mapped onto unused cp1252 code points and declared in the font's /Differences
# array, so text extractors resolve them back to the right Unicode character.
_EXTRA_GLYPHS = {
    "\u0150": (0x81, "Ohungarumlaut"),
    "\u0151": (0x8D, "ohungarumlaut"),
    "\u0170": (0x8F, "Uhungarumlaut"),
    "\u0171": (0x90, "uhungarumlaut"),
}


def _pdf_escape(text: str) -> bytes:
    # NOTE: the substitution must happen at BYTE level. The chosen code points
    # (0x81/0x8D/0x8F/0x90) are exactly the ones cp1252 leaves undefined, so
    # encoding a chr(0x8D) through cp1252 would turn it into '?' again.
    out = bytearray()
    for char in text:
        if char in _EXTRA_GLYPHS:
            raw = bytes([_EXTRA_GLYPHS[char][0]])
        else:
            raw = char.encode("cp1252", "replace")
        for byte in raw:
            if byte in (0x28, 0x29, 0x5C):  # ( ) \
                out.append(0x5C)
            out.append(byte)
    return bytes(out)


def _layout_lines(blocks: list[tuple[str, str]]) -> list[list[tuple[str, str]]]:
    """Flow (kind, text) blocks into pages of (kind, line) tuples."""
    flowed: list[tuple[str, str]] = []
    for kind, text in blocks:
        if kind == "pagebreak":
            flowed.append(("pagebreak", ""))
            continue
        if not text:
            flowed.append(("", ""))
            continue
        width = 74 if kind in ("h1", "h2") else _WRAP
        for line in textwrap.wrap(text, width=width) or [""]:
            flowed.append((kind, line))
        if kind in ("h1", "h2"):
            flowed.append(("", ""))
    pages: list[list[tuple[str, str]]] = [[]]
    for item in flowed:
        if item[0] == "pagebreak" or len(pages[-1]) >= _LINES_PER_PAGE:
            if pages[-1]:
                pages.append([])
            if item[0] == "pagebreak":
                continue
        pages[-1].append(item)
    return [p for p in pages if p] or [[("", "")]]


def _pdf_page_stream(lines: list[tuple[str, str]], page_no: int, total: int, footer: str) -> bytes:
    parts = [b"BT\n"]
    y = _MARGIN_TOP
    for kind, text in lines:
        if kind == "h1":
            font, size = b"/F2", 15
        elif kind == "h2":
            font, size = b"/F2", 12
        elif kind == "mono":
            font, size = b"/F3", 9
        else:
            font, size = b"/F1", 10
        parts.append(b"%s %d Tf\n" % (font, size))
        parts.append(b"1 0 0 1 %d %d Tm\n" % (_MARGIN_X, y))
        parts.append(b"(%s) Tj\n" % _pdf_escape(text))
        y -= _LEADING
    parts.append(b"/F1 8 Tf\n1 0 0 1 %d %d Tm\n(%s) Tj\n"
                 % (_MARGIN_X, 40, _pdf_escape("%s  |  page %d of %d" % (footer, page_no, total))))
    parts.append(b"ET\n")
    return b"".join(parts)


def _build_pdf(objects: list[bytes]) -> bytes:
    out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + obj + b"\nendobj\n"
    xref_at = len(out)
    out += b"xref\n0 %d\n" % (len(objects) + 1)
    out += b"0000000000 65535 f \n"
    for off in offsets[1:]:
        out += b"%010d 00000 n \n" % off
    out += (b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
            % (len(objects) + 1, xref_at))
    return bytes(out)


def write_pdf(path: str, title: str, blocks: list[tuple[str, str]], footer: str = "") -> int:
    """blocks: list of (kind, text); kind in 'h1' | 'h2' | 'mono' | '' | 'pagebreak'."""
    os.makedirs(os.path.dirname(path), exist_ok=True)
    pages = _layout_lines([("h1", title)] + list(blocks))
    total = len(pages)
    footer = footer or title

    n_pages = len(pages)
    # object ids: 1 catalog, 2 pages, then page objects, contents, the shared
    # encoding and the three fonts
    page_ids = list(range(3, 3 + n_pages))
    content_ids = list(range(3 + n_pages, 3 + 2 * n_pages))
    encoding_id = 3 + 2 * n_pages
    font_ids = [encoding_id + 1, encoding_id + 2, encoding_id + 3]

    objects: list[bytes] = []
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = b" ".join(b"%d 0 R" % pid for pid in page_ids)
    objects.append(b"<< /Type /Pages /Count %d /Kids [%s] >>" % (n_pages, kids))
    for idx, pid in enumerate(page_ids):
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 %d %d] "
            b"/Resources << /Font << /F1 %d 0 R /F2 %d 0 R /F3 %d 0 R >> >> "
            b"/Contents %d 0 R >>"
            % (_PAGE_W, _PAGE_H, font_ids[0], font_ids[1], font_ids[2], content_ids[idx])
        )
    for idx, lines in enumerate(pages):
        stream = _pdf_page_stream(lines, idx + 1, total, footer)
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"endstream")
    differences = b" ".join(
        b"%d /%s" % (code, name.encode("ascii"))
        for code, name in sorted(_EXTRA_GLYPHS.values())
    )
    objects.append(b"<< /Type /Encoding /BaseEncoding /WinAnsiEncoding /Differences [%s] >>"
                   % differences)
    for base in (b"/Helvetica", b"/Helvetica-Bold", b"/Courier"):
        objects.append(
            b"<< /Type /Font /Subtype /Type1 /BaseFont %s /Encoding %d 0 R >>"
            % (base, encoding_id)
        )

    with open(path, "wb") as f:
        f.write(_build_pdf(objects))
    return os.path.getsize(path)


def write_scanned_pdf(path: str, pages: int = 2, seed: int = 7) -> int:
    """Image-only PDF: a synthetic 'scan' raster with NO text operators.

    MarkItDown (and any non-OCR extractor) yields an empty string for this file,
    which is exactly the `extraction_empty` case the demo needs to exercise.
    """
    os.makedirs(os.path.dirname(path), exist_ok=True)
    rng = random.Random(seed)
    w, h = 620, 877  # ~75 dpi A4, grayscale

    def raster() -> bytes:
        rows = bytearray()
        for y in range(h):
            row = bytearray(b"\xf2" * w)  # off-white paper
            # fake scanned text lines
            if 90 < y < 800 and (y % 26) < 7 and rng.random() < 0.85:
                x = 70
                while x < w - 90:
                    word = rng.randint(14, 58)
                    for i in range(min(word, w - 90 - x)):
                        row[x + i] = rng.randint(20, 90)
                    x += word + rng.randint(6, 14)
            # scanner edge shadow
            for i in range(6):
                row[i] = 0xB0 - i * 6
            rows += row
        return bytes(rows)

    n_pages = max(1, pages)
    page_ids = list(range(3, 3 + n_pages))
    content_ids = list(range(3 + n_pages, 3 + 2 * n_pages))
    image_ids = list(range(3 + 2 * n_pages, 3 + 3 * n_pages))

    objects: list[bytes] = []
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = b" ".join(b"%d 0 R" % pid for pid in page_ids)
    objects.append(b"<< /Type /Pages /Count %d /Kids [%s] >>" % (n_pages, kids))
    for idx, _ in enumerate(page_ids):
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 %d %d] "
            b"/Resources << /XObject << /Im0 %d 0 R >> >> /Contents %d 0 R >>"
            % (_PAGE_W, _PAGE_H, image_ids[idx], content_ids[idx])
        )
    for _ in range(n_pages):
        stream = b"q %d 0 0 %d 0 0 cm /Im0 Do Q\n" % (_PAGE_W, _PAGE_H)
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"endstream")
    for _ in range(n_pages):
        data = zlib.compress(raster(), 6)
        objects.append(
            b"<< /Type /XObject /Subtype /Image /Width %d /Height %d /ColorSpace /DeviceGray "
            b"/BitsPerComponent 8 /Filter /FlateDecode /Length %d >>\nstream\n" % (w, h, len(data))
            + data
            + b"\nendstream"
        )

    with open(path, "wb") as f:
        f.write(_build_pdf(objects))
    return os.path.getsize(path)
